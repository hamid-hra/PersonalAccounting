"""
آداپتور صورتحساب اکسل بانک سامان.

ساختار فایل (تأییدشده روی نمونهٔ واقعی):
  • سطرهای ۱ تا ۱۰: بلوک سرصفحه، برچسب‌ها به‌صورت «کلید: مقدار» در سلول‌های پراکنده
  • سطر هدر جدول: سلول‌هایی با متن «ردیف»، «تاریخ»، «شرح سند»، «مانده» …
  • سطرهای بعدی: تراکنش‌ها، از جدید به قدیم
  • ستون‌های خالی زیاد بین ستون‌های واقعی — پس ستون‌ها باید از روی متن هدر
    پیدا شوند، نه با حرف ثابت.
"""

import re
from pathlib import Path

from openpyxl import load_workbook

from app.importers.base import BankImporter, ParsedStatement, RawTxn, StatementMeta
from app.importers.registry import register
from app.enums import Bank
from app.services.extract import clean_amount
from app.services.normalize import fa_digits, normalize

# سرصفحه‌ها؛ ترتیب مهم است — «مانده از قبل» باید قبل از «مانده» بررسی شود
_META_LABELS: list[tuple[str, str]] = [
    ("مانده از قبل", "opening_balance"),
    ("جمع کل واریز", "total_deposit"),
    ("جمع کل برداشت", "total_withdraw"),
    ("معدل موجودی", "average_balance"),
    ("نام صاحب حساب", "owner_name"),
    ("تاریخ افتتاح حساب", "opened_at"),
    ("شماره شبا", "iban"),
    ("نوع ارز", "currency"),
    ("مانده", "closing_balance"),
]

_PERIOD_RE = re.compile(
    r"از\s*تاریخ\s*(?P<from>[\d/]+(?:\s+[\d:]+)?)\s*تا\s*تاریخ\s*(?P<to>[\d/]+(?:\s+[\d:]+)?)"
)
_IBAN_RE = re.compile(r"(IR\d{24})", re.IGNORECASE)

# نگاشت متن هدر ستون → نام منطقی
_COLUMNS: list[tuple[str, str]] = [
    ("شرح سند", "description"),
    ("شماره سند", "doc_number"),
    ("نوع تراکنش", "tx_type"),
    ("برداشت", "withdraw"),
    ("واریز", "deposit"),
    ("مانده", "balance"),
    ("تاریخ", "date"),
    ("ردیف", "row_index"),
]

_HEADER_REQUIRED = {"row_index", "date", "balance", "description"}


def _label_of(text: str) -> str:
    """«مانده (ریال)» → «مانده»"""
    return re.sub(r"\s*\(.*?\)\s*", "", normalize(text)).strip()


def _dmy_to_jalali(value: str) -> str | None:
    """دورهٔ صورتحساب به‌صورت «10/06/1404 00:00» (روز/ماه/سال) نوشته می‌شود."""
    m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", fa_digits(value or ""))
    if not m:
        return None
    d, mo, y = m.groups()
    return f"{int(y):04d}/{int(mo):02d}/{int(d):02d}"


class SamanImporter(BankImporter):
    bank = Bank.SAMAN
    label_fa = "بانک سامان"

    # ---------- تشخیص ----------
    def sniff(self, path: Path) -> bool:
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            return False
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            blob = " ".join(
                normalize(str(c))
                for row in ws.iter_rows(min_row=1, max_row=20, values_only=True)
                for c in row
                if c is not None
            )
        finally:
            wb.close()
        return "شرح سند" in blob and "نوع تراکنش" in blob and "ردیف" in blob

    # ---------- پارس ----------
    def parse(self, path: Path) -> ParsedStatement:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()

        warnings: list[str] = []
        header_idx, colmap = self._find_header(rows)
        if header_idx is None:
            raise ValueError("سطر هدر جدول تراکنش‌ها در فایل پیدا نشد.")

        meta = self._parse_meta(rows[:header_idx])
        txns = self._parse_rows(rows[header_idx + 1 :], colmap, warnings)
        return ParsedStatement(meta=meta, transactions=txns, warnings=warnings)

    # ---------- بخش‌های داخلی ----------
    def _find_header(self, rows: list[list]) -> tuple[int | None, dict[str, int]]:
        for idx, row in enumerate(rows[:40]):
            colmap: dict[str, int] = {}
            for col, cell in enumerate(row):
                if cell is None:
                    continue
                label = _label_of(str(cell))
                for needle, key in _COLUMNS:
                    if key not in colmap and label == needle:
                        colmap[key] = col
                        break
            if _HEADER_REQUIRED.issubset(colmap):
                return idx, colmap
        return None, {}

    def _parse_meta(self, header_rows: list[list]) -> StatementMeta:
        meta = StatementMeta(bank=self.bank)
        found: dict[str, str] = {}

        for row in header_rows:
            for cell in row:
                if cell is None:
                    continue
                text = normalize(str(cell))
                if not text:
                    continue

                if m := _PERIOD_RE.search(text):
                    meta.period_from_jalali = _dmy_to_jalali(m.group("from"))
                    meta.period_to_jalali = _dmy_to_jalali(m.group("to"))

                for needle, key in _META_LABELS:
                    if key in found:
                        continue
                    m2 = re.match(rf"^{re.escape(needle)}\s*[:：]\s*(.+)$", text)
                    if m2:
                        found[key] = m2.group(1).strip()
                        break

        meta.owner_name = found.get("owner_name")
        meta.currency = found.get("currency")
        meta.opened_at_jalali = self._clean_jalali(found.get("opened_at"))

        if iban_raw := found.get("iban"):
            if m := _IBAN_RE.search(fa_digits(iban_raw).replace(" ", "")):
                meta.iban = m.group(1).upper()
                # شبای ایران: IR + ۲ رقم کنترل + ۳ رقم کد بانک + ۱۹ رقم شماره حساب
                meta.account_number = meta.iban[7:].lstrip("0") or meta.iban[7:]

        for key, attr in [
            ("opening_balance", "opening_balance_rial"),
            ("closing_balance", "closing_balance_rial"),
            ("total_deposit", "total_deposit_rial"),
            ("total_withdraw", "total_withdraw_rial"),
        ]:
            if key in found:
                setattr(meta, attr, clean_amount(found[key]))

        return meta

    @staticmethod
    def _clean_jalali(value: str | None) -> str | None:
        if not value:
            return None
        m = re.search(r"(1[34]\d{2})/(\d{1,2})/(\d{1,2})", fa_digits(value))
        return f"{int(m[1]):04d}/{int(m[2]):02d}/{int(m[3]):02d}" if m else None

    def _parse_rows(
        self, rows: list[list], colmap: dict[str, int], warnings: list[str]
    ) -> list[RawTxn]:
        def cell(row: list, key: str):
            col = colmap.get(key)
            if col is None or col >= len(row):
                return None
            return row[col]

        out: list[RawTxn] = []
        for row in rows:
            raw_date = cell(row, "date")
            raw_seq = cell(row, "row_index")
            if raw_date is None or raw_seq is None:
                continue

            date_text = fa_digits(str(raw_date)).strip()
            if not re.match(r"^1[34]\d{2}[/\-]\d{1,2}[/\-]\d{1,2}", date_text):
                continue  # سطرهای جمع/پانویس

            description = str(cell(row, "description") or "").strip()
            balance = cell(row, "balance")
            try:
                seq = int(float(fa_digits(str(raw_seq))))
            except ValueError:
                seq = None

            out.append(
                RawTxn(
                    jalali_datetime=self._normalize_dt(date_text),
                    withdraw_rial=clean_amount(cell(row, "withdraw")),
                    deposit_rial=clean_amount(cell(row, "deposit")),
                    balance_after_rial=clean_amount(balance) if balance is not None else None,
                    description_raw=description,
                    bank_tx_type=(str(cell(row, "tx_type")).strip() or None)
                    if cell(row, "tx_type")
                    else None,
                    doc_number=(str(cell(row, "doc_number")).strip() or None)
                    if cell(row, "doc_number")
                    else None,
                    row_index=seq,
                )
            )

        if not out:
            warnings.append("هیچ سطر تراکنشی در فایل پیدا نشد.")
        return out

    @staticmethod
    def _normalize_dt(text: str) -> str:
        m = re.match(
            r"^(1[34]\d{2})[/\-](\d{1,2})[/\-](\d{1,2})"
            r"(?:[ T]+(\d{1,2}):(\d{2})(?::(\d{2}))?)?",
            text,
        )
        if not m:
            return text
        y, mo, d = int(m[1]), int(m[2]), int(m[3])
        hh, mi, ss = int(m[4] or 0), int(m[5] or 0), int(m[6] or 0)
        return f"{y:04d}/{mo:02d}/{d:02d} {hh:02d}:{mi:02d}:{ss:02d}"


register(SamanImporter())
